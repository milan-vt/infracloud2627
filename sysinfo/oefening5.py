from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_BESTAND = "sysinfo_computers.xlsx"


def schrijf_xlsx(rijen, bestand=XLSX_BESTAND):
    wb = Workbook()
    ws = wb.active
    ws.title = "Computers"

    # Gegevens: zelfde kolommen als de CSV, één rij per computer.
    # None wordt een lege cel; getallen blijven echte getallen.
    ws.append(KOLOMMEN)
    for rij in rijen:
        ws.append([rij.get(k) for k in KOLOMMEN])

    # Opmaak: dit kan CSV niet
    for rij in ws.iter_rows():
        for cel in rij:
            cel.font = Font(name="Arial")
    for cel in ws[1]:  # kopregel
        cel.font = Font(name="Arial", bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="305496")
    for kolom in ("ram_totaal_gib", "opslag_capaciteit_gib"):
        letter = get_column_letter(KOLOMMEN.index(kolom) + 1)
        for cel in ws[letter][1:]:
            cel.number_format = "0.00"  # Excel toont zelf , of . volgens de taalinstelling
    for i in range(1, len(KOLOMMEN) + 1):
        letter = get_column_letter(i)
        breedte = max(len(str(c.value or "")) for c in ws[letter])
        ws.column_dimensions[letter].width = breedte + 2
    ws.freeze_panes = "A2"               # kopregel blijft zichtbaar bij scrollen
    ws.auto_filter.ref = ws.dimensions   # filterknoppen op de kopregel

    wb.save(bestand)


def lees_xlsx(bestand=XLSX_BESTAND):
    """Leest het werkblad in als een lijst van dicts, net zoals lees_csv()."""
    ws = load_workbook(bestand).active
    rijen = list(ws.iter_rows(values_only=True))
    kop = rijen[0]
    return [dict(zip(kop, rij)) for rij in rijen[1:]]


# CSV inlezen (taak 4) -> spreadsheetbestand schrijven -> controleren
rijen = lees_csv()
schrijf_xlsx(rijen)
print("Gegevens in XLSX gelijk aan CSV:", lees_xlsx() == rijen)